#!/usr/bin/env python3
"""board-slide.py — rewrite one slide of the quarterly board deck as a stand-alone Google Slides file.

WHY THIS EXISTS
---------------
Every quarter the board deck (Google Slides, owned by the CEO's office) is a copy of the previous
one, and the marketing slide "Marketing Efforts" still carries last quarter's text. Vadim fills it
as a separate one-slide file and imports it into the deck himself (File -> Import slides). Slides
API is closed through the oo proxy, so the path is: export the deck as .pptx -> cut out the one
slide with python-pptx (theme, fonts, logo and divider stay native) -> replace the text of its
text boxes -> upload as a new Google Slides file via the Drive multipart upload.

THE CONTENT FILE IS NOT IN GIT
  The slide text carries board numbers (deals, targets, plans) and this repo is PUBLIC. Keep the
  JSON in the scratchpad or another private place and pass its path. Format:

  {
    "slide_index": 15,              # 0-based slide in the exported deck
    "remove_shapes": [235],         # shape ids to drop (e.g. last quarter's table)
    "match_height": {"234": 233},   # give shape 234 the height of shape 233
    "font": "Lato", "size_pt": 10, "accent": "143DFF",
    "boxes": {                      # shape id -> paragraphs; all old paragraphs are replaced
      "233": [
        ["head", [["Q3'2026", "blue"], [" Updates: ", "b"], ["Headline.", "b"]]],
        ["item", [["Bold lead", "b"], [" — plain text.", ""]]],
        ["note", [["Missed: ", "b"], ["plain text.", ""]]],
        ["note_tight", [["Results: ", "b"], ["plain text.", ""]]]
      ]
    }
  }

  Paragraph kinds: head (no space before), item (auto-numbered 1., 2., ...), note (plain, a gap
  before), note_tight (plain, a smaller gap). Run styles: "" plain, "b" bold, "blue" bold accent.
  Shape ids come from `inspect`.

USAGE (python-pptx lives only in the Hermes venv)
  PY=~/.hermes/hermes-agent/venv/bin/python
  oo connector run googledrive --action files.export \
      --data '{"fileId":"<deck id>","mimeType":"application/vnd.openxmlformats-officedocument.presentationml.presentation"}'
  curl -o deck.pptx "<transitUrl>"
  $PY board-slide.py inspect deck.pptx 15
  $PY board-slide.py build deck.pptx content.json slide.pptx
  $PY board-slide.py upload slide.pptx "Marketing Efforts Q3 2026 (board slide)" [<file id to update>]
  $PY board-slide.py render <file id> out-prefix      # PDF -> PNG to check the layout

  `upload` without a file id creates a new, unshared file in My Drive of the oo Google account;
  with an id it replaces that file's content (same link). `render` exports through Drive, because
  the proxy's /export returns the PDF as broken text.
"""
import base64
import json
import os
import subprocess
import sys
import tempfile
import urllib.request

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SPACE_BEFORE = {"head": 0, "note": 1000, "note_tight": 600}  # hundredths of a point
ITEM_FIRST, ITEM_NEXT = 1000, 500


def q(name):
    return "{%s}%s" % (A, name)


def run_props(etree, style, cfg, end=False):
    el = etree.Element(q("endParaRPr" if end else "rPr"))
    if style in ("b", "blue"):
        el.set("b", "1")
    el.set("lang", "en")
    el.set("sz", str(int(cfg["size_pt"] * 100)))
    fill = etree.SubElement(el, q("solidFill"))
    if style == "blue":
        etree.SubElement(fill, q("srgbClr")).set("val", cfg["accent"])
    else:
        etree.SubElement(fill, q("schemeClr")).set("val", "dk1")
    for tag in ("latin", "ea", "cs", "sym"):
        etree.SubElement(el, q(tag)).set("typeface", cfg["font"])
    return el


def para_props(etree, kind, first_item, cfg):
    el = etree.Element(q("pPr"))
    numbered = kind == "item"
    el.set("indent", "-228600" if numbered else "0")
    el.set("marL", "228600" if numbered else "0")
    el.set("lvl", "0")
    el.set("rtl", "0")
    el.set("algn", "l")
    etree.SubElement(etree.SubElement(el, q("lnSpc")), q("spcPct")).set("val", "115000")
    before = (ITEM_FIRST if first_item else ITEM_NEXT) if numbered else SPACE_BEFORE[kind]
    etree.SubElement(etree.SubElement(el, q("spcBef")), q("spcPts")).set("val", str(before))
    etree.SubElement(etree.SubElement(el, q("spcAft")), q("spcPts")).set("val", "0")
    if numbered:
        etree.SubElement(etree.SubElement(el, q("buClr")), q("schemeClr")).set("val", "dk1")
        etree.SubElement(el, q("buSzPts")).set("val", str(int(cfg["size_pt"] * 100)))
        etree.SubElement(el, q("buFont")).set("typeface", cfg["font"])
        etree.SubElement(el, q("buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(el, q("buNone"))
    return el


def fill_box(etree, shape, paragraphs, cfg):
    body = shape._element.find(".//{%s}txBody" % P)
    for p in body.findall(q("p")):
        body.remove(p)
    seen_item = False
    for kind, runs in paragraphs:
        if kind != "item" and kind not in SPACE_BEFORE:
            sys.exit(f"unknown paragraph kind: {kind}")
        p = etree.SubElement(body, q("p"))
        p.append(para_props(etree, kind, kind == "item" and not seen_item, cfg))
        seen_item = seen_item or kind == "item"
        for text, style in runs:
            r = etree.SubElement(p, q("r"))
            r.append(run_props(etree, style, cfg))
            etree.SubElement(r, q("t")).text = text
        p.append(run_props(etree, runs[-1][1], cfg, end=True))


def keep_one_slide(prs, index):
    ids = prs.slides._sldIdLst
    for i, sld_id in reversed(list(enumerate(list(ids)))):
        if i != index:
            prs.part.drop_rel(sld_id.get("{%s}id" % R))
            ids.remove(sld_id)
    return prs.slides[0]


def cmd_inspect(deck, index):
    from pptx import Presentation

    slide = Presentation(deck).slides[int(index)]
    for sh in slide.shapes:
        text = sh.text_frame.text.replace("\n", " / ")[:90] if sh.has_text_frame else ""
        print(sh.shape_id, sh.shape_type, f"x={sh.left / 12700:.0f} y={sh.top / 12700:.0f}",
              f"w={sh.width / 12700:.0f} h={sh.height / 12700:.0f}", text)


def cmd_build(deck, content_path, out):
    from lxml import etree
    from pptx import Presentation

    content = json.load(open(content_path, encoding="utf-8"))
    cfg = {"font": content.get("font", "Lato"), "size_pt": content.get("size_pt", 10),
           "accent": content.get("accent", "143DFF")}
    prs = Presentation(deck)
    slide = keep_one_slide(prs, content["slide_index"])
    shapes = {str(sh.shape_id): sh for sh in slide.shapes}
    for sid in content.get("remove_shapes", []):
        el = shapes[str(sid)]._element
        el.getparent().remove(el)
    for sid, paragraphs in content["boxes"].items():
        fill_box(etree, shapes[sid], paragraphs, cfg)
    for target, source in content.get("match_height", {}).items():
        shapes[target].height = shapes[str(source)].height
    prs.save(out)
    print("saved", out)


def cmd_upload(pptx, name, file_id=None):
    boundary = "boardslide7f3a"
    meta = {"name": name} if file_id else {"name": name, "mimeType": "application/vnd.google-apps.presentation"}
    data = base64.b64encode(open(pptx, "rb").read()).decode()
    body = (f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{json.dumps(meta)}\r\n"
            f"--{boundary}\r\nContent-Type: application/vnd.openxmlformats-officedocument.presentationml.presentation\r\n"
            f"Content-Transfer-Encoding: base64\r\n\r\n{data}\r\n--{boundary}--")
    url = ("https://www.googleapis.com/upload/drive/v3/files" + (f"/{file_id}" if file_id else "")
           + "?uploadType=multipart&fields=id,name,webViewLink")
    req = {"method": "PATCH" if file_id else "POST", "endpoint": url,
           "headers": {"Content-Type": f"multipart/related; boundary={boundary}"}, "body": body}
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as f:
        json.dump(req, f)
    try:
        r = subprocess.run(["oo", "connector", "proxy", "googledrive", "-d", "@" + f.name, "--json"],
                           capture_output=True, text=True)
    finally:
        os.unlink(f.name)
    try:
        resp = json.loads(r.stdout)["data"]
        print(resp["status"], json.dumps(resp["data"]))
    except (ValueError, KeyError):
        sys.exit(f"upload failed:\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}")


def cmd_render(file_id, prefix):
    r = subprocess.run(["oo", "connector", "run", "googledrive", "--action", "files.export", "--data",
                        json.dumps({"fileId": file_id, "mimeType": "application/pdf"})],
                       capture_output=True, text=True)
    out = r.stdout
    try:
        url = json.loads(out[out.index("{"):])["transitUrl"]
    except (ValueError, KeyError):
        sys.exit(f"export failed:\n{out[-1500:]}\n{r.stderr[-1500:]}")
    pdf = prefix + ".pdf"
    urllib.request.urlretrieve(url, pdf)
    subprocess.run(["pdftoppm", "-r", "130", "-png", pdf, prefix], check=True)
    print("rendered", pdf, "->", prefix + "-*.png")


COMMANDS = {"inspect": (cmd_inspect, 2), "build": (cmd_build, 3), "upload": (cmd_upload, 2),
            "render": (cmd_render, 2)}

if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in COMMANDS:
        sys.exit(__doc__)
    fn, required = COMMANDS[sys.argv[1]]
    args = sys.argv[2:]
    if len(args) < required:
        sys.exit(__doc__)
    fn(*args)
